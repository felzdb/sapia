
from io import BytesIO
from pathlib import Path

import fitz
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt


SECTION_TITLES = {
    "QUALIFICAÇÃO",
    "DOS FATOS",
    "DO DIREITO",
    "DOS PEDIDOS",
    "PARTICULARIDADES DO CASO",
    "DO VALOR DA CAUSA",
}


def is_heading(text: str, first_paragraph: bool = False) -> bool:
    """Identifica os títulos da petição."""
    return (
        first_paragraph
        or text in SECTION_TITLES
        or text.startswith("PETIÇÃO INICIAL")
    )


def export_petition_docx(content: str) -> bytes:
    """Gera um documento Word a partir do texto da petição."""

    document = Document()
    section = document.sections[0]

    # Página A4 e margens
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(3)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(2)

    # Formatação padrão
    normal_style = document.styles["Normal"]
    normal_style.font.name = "Arial"
    normal_style.font.size = Pt(12)
    normal_style.paragraph_format.line_spacing = 1.5
    normal_style.paragraph_format.space_after = Pt(6)

    paragraphs = [
        line.strip()
        for line in content.splitlines()
        if line.strip()
    ]

    for index, text in enumerate(paragraphs):
        paragraph = document.add_paragraph()

        if is_heading(text, index == 0):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_with_next = True

            run = paragraph.add_run(text)
            run.bold = True

        else:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            paragraph.paragraph_format.first_line_indent = Cm(1.25)
            paragraph.add_run(text)

    # Geração em memória
    buffer = BytesIO()
    document.save(buffer)

    return buffer.getvalue()


def export_petition_pdf(content: str) -> bytes:
    """Gera PDF com paginação automática e fontes portáteis."""

    fonts_directory = Path(__file__).resolve().parent / "fonts"

    font_regular = fonts_directory / "LiberationSans-Regular.ttf"
    font_bold = fonts_directory / "LiberationSans-Bold.ttf"

    if not font_regular.is_file() or not font_bold.is_file():
        raise FileNotFoundError(
            "Fontes Liberation Sans não encontradas na pasta app/fonts."
        )

    regular_font = fitz.Font(fontfile=str(font_regular))
    bold_font = fitz.Font(fontfile=str(font_bold))

    # Dimensões da página A4 em pontos
    page_width = 595.28
    page_height = 841.89

    # Margens: 3 cm superior/esquerda e 2 cm inferior/direita
    left_margin = 85.04
    right_margin = 56.69
    top_margin = 85.04
    bottom_margin = 56.69

    available_width = page_width - left_margin - right_margin
    bottom_limit = page_height - bottom_margin

    font_size = 12
    line_height = 18
    paragraph_spacing = 12
    first_line_indent = 35.43  # 1,25 cm

    pdf = fitz.open()

    try:
        def new_page():
            page = pdf.new_page(
                width=page_width,
                height=page_height,
            )

            page.insert_font(
                fontname="LiberationRegular",
                fontfile=str(font_regular),
            )

            page.insert_font(
                fontname="LiberationBold",
                fontfile=str(font_bold),
            )

            return page

        def wrap_text(text, font, indent=0):
            """Quebra linhas respeitando largura e recuo inicial."""

            words = text.split()
            lines = []
            current_line = ""

            for word in words:
                candidate = (
                    f"{current_line} {word}".strip()
                )

                max_width = available_width

                if not lines:
                    max_width -= indent

                if font.text_length(
                    candidate,
                    fontsize=font_size,
                ) <= max_width:
                    current_line = candidate
                    continue

                if current_line:
                    lines.append(current_line)
                    current_line = ""

                # Trata palavras maiores que a largura da página
                max_width = available_width
                if not lines:
                    max_width -= indent

                for character in word:
                    candidate = current_line + character

                    if (
                        font.text_length(
                            candidate,
                            fontsize=font_size,
                        ) > max_width
                        and current_line
                    ):
                        lines.append(current_line)
                        current_line = character
                        max_width = available_width
                    else:
                        current_line = candidate

            if current_line:
                lines.append(current_line)

            return lines

        paragraphs = [
            line.strip()
            for line in content.splitlines()
            if line.strip()
        ]

        page = new_page()
        y = top_margin

        for index, paragraph in enumerate(paragraphs):
            heading = is_heading(
                paragraph,
                first_paragraph=(index == 0),
            )

            font = bold_font if heading else regular_font
            font_name = (
                "LiberationBold"
                if heading
                else "LiberationRegular"
            )

            indent = 0 if heading else first_line_indent

            lines = wrap_text(
                paragraph,
                font,
                indent,
            )

            for line_index, line in enumerate(lines):
                if y + line_height > bottom_limit:
                    page = new_page()
                    y = top_margin

                if heading:
                    text_width = font.text_length(
                        line,
                        fontsize=font_size,
                    )

                    x = left_margin + (
                        available_width - text_width
                    ) / 2

                else:
                    x = left_margin

                    if line_index == 0:
                        x += first_line_indent

                                # Justifica as linhas intermediárias dos parágrafos.
                # A última linha permanece alinhada à esquerda.
                should_justify = (
                    not heading
                    and line_index < len(lines) - 1
                    and " " in line
                )

                if should_justify:
                    words = line.split()
                    text_width = sum(
                        font.text_length(word, fontsize=font_size)
                        for word in words
                    )

                    line_width = available_width
                    if line_index == 0:
                        line_width -= first_line_indent

                    space_width = (
                        line_width - text_width
                    ) / (len(words) - 1)

                    current_x = x

                    for word in words:
                        page.insert_text(
                            (current_x, y + font_size),
                            word,
                            fontname=font_name,
                            fontsize=font_size,
                        )

                        current_x += (
                            font.text_length(
                                word,
                                fontsize=font_size,
                            )
                            + space_width
                        )
                else:
                    page.insert_text(
                        (x, y + font_size),
                        line,
                        fontname=font_name,
                        fontsize=font_size,
                    )

                y += line_height

            y += paragraph_spacing

        return pdf.tobytes()

    finally:
        pdf.close()
